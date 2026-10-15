from flask import Flask, render_template, request
import pandas as pd
from datetime import datetime
import random
import os
from openpyxl.styles import Font, Alignment
from flask import send_file

from openpyxl import load_workbook
from werkzeug.utils import secure_filename

app = Flask(__name__)

app.config['MAX_CONTENT_LENGTH'] = 2 * 1024 * 1024 * 1024

ARCHIVO_EXCEL = r"C:\Users\Usuario\OneDrive - ERCO ENERGIA SAS\Reporte informes mantenimientos.xlsx"

EXTENSIONES_PERMITIDAS = {
    "jpg",
    "jpeg",
    "png",
    "heic",
    "heif",
    "webp",
    "bmp",
    "gif",
    "tif",
    "tiff",
    "pdf"
}

@app.route("/")
def menu():
    return render_template("menu.html")

RUTA_PLANOS = r"C:\Users\Usuario\OneDrive - ERCO ENERGIA SAS\AOM - 001_AOM_COMERCIAL&INDUSTRIAL\02_Ejecucion\2.4_AOM_Centro\2.4.4_Mantenimiento\2.4.4.3_Anexos\Planos"

@app.route("/planos")
def planos():

    df_proyectos = pd.read_excel(
        ARCHIVO_EXCEL,
        sheet_name="OMCE",
        engine="openpyxl"
    )

    proyectos = (
        df_proyectos["Proyectos"]
        .dropna()
        .tolist()
    )

    return render_template(
        "planos.html",
        proyectos=proyectos,
        archivos=[],
        filtro=""
    )

@app.route("/buscar_planos")
def buscar_planos():

    proyecto = request.args.get("proyecto", "").strip()

    resultados = []

    print("RUTA BASE:")
    print(RUTA_PLANOS)

    for raiz, carpetas, archivos in os.walk(RUTA_PLANOS):

        print("CARPETA:", raiz)

        for archivo in archivos:

            print("ARCHIVO:", archivo)

            if archivo.lower().endswith(".pdf"):

                resultados.append({
                    "nombre": archivo,
                    "ruta": os.path.join(raiz, archivo)
                })

    print("TOTAL PDF:", len(resultados))

    df_proyectos = pd.read_excel(
        ARCHIVO_EXCEL,
        sheet_name="OMCE",
        engine="openpyxl"
    )

    proyectos = (
        df_proyectos["Proyectos"]
        .dropna()
        .tolist()
    )

    return render_template(
        "planos.html",
        proyectos=proyectos,
        archivos=resultados,
        filtro=proyecto
    )

@app.route("/abrir_plano")
def abrir_plano():

    ruta = request.args.get("ruta")

    return send_file(
        ruta,
        mimetype="application/pdf"
    )


@app.route("/consultar")
def consultar():

    df = pd.read_excel(
        ARCHIVO_EXCEL,
        sheet_name="Informe OMCE",
        engine="openpyxl"
    )

    filtro_id = request.args.get("id", "")
    filtro_proyecto = request.args.get("proyecto", "")
    filtro_lider = request.args.get("lider", "")

    if filtro_id:
        df = df[
            df["ID"].astype(str).str.contains(
                filtro_id,
                case=False,
                na=False
            )
        ]

    if filtro_proyecto:
        df = df[
            df["Proyecto"].astype(str).str.contains(
                filtro_proyecto,
                case=False,
                na=False
            )
        ]

    if filtro_lider:
        df = df[
            df["LiderEncargado"].astype(str).str.contains(
                filtro_lider,
                case=False,
                na=False
            )
        ]

    registros = df.to_dict(
        orient="records"
    )

    return render_template(
        "consultar.html",
        registros=registros,
        filtro_id=filtro_id,
        filtro_proyecto=filtro_proyecto,
        filtro_lider=filtro_lider
    )


@app.route("/mantenimiento", methods=["GET", "POST"])
def mantenimiento():

    print("ENTRO A MANTENIMIENTO")
    
    df_proyectos = pd.read_excel(
        ARCHIVO_EXCEL,
        sheet_name="OMCE",
        engine="openpyxl"
    )

    proyectos = df_proyectos["Proyectos"].dropna().tolist()

    if request.method == "POST":

        proyecto = request.form.get("proyecto")
        lider = request.form.get("lider")
        fecha = request.form.get("fecha")
        tipo = request.form.get("tipomantenimiento")
        servicio = request.form.get("servicio")
        novedades = request.form.get("novedades")
        pendientes = request.form.get("pendientes")

        if not proyecto or not lider or not fecha or not tipo:

            return """
            <html>
            <body style="font-family:Arial">

            <h2>Error</h2>

            <p>Debe diligenciar todos los campos obligatorios.</p>

            <button onclick="window.history.back()">
            Volver
            </button>

            </body>
            </html>
            """

        fecha_actual = datetime.now()

        id_mtto = (
            "MTTO-"
            + fecha_actual.strftime("%Y%m%d")
            + "-"
            + str(random.randint(100000, 999999))
        )

        meses = {
            1: "01. Enero",
            2: "02. Febrero",
            3: "03. Marzo",
            4: "04. Abril",
            5: "05. Mayo",
            6: "06. Junio",
            7: "07. Julio",
            8: "08. Agosto",
            9: "09. Septiembre",
            10: "10. Octubre",
            11: "11. Noviembre",
            12: "12. Diciembre"
        }

        ano = str(fecha_actual.year)
        mes = meses[fecha_actual.month]

        RUTA_SHAREPOINT_MTTO = r"C:\Users\Usuario\OneDrive - ERCO ENERGIA SAS\AOM - 001_AOM_COMERCIAL&INDUSTRIAL\02_Ejecucion\2.4_AOM_Centro\2.4.4_Mantenimiento\2.4.4.3_Anexos\01. Mantenimientos"

        carpeta_mtto = os.path.join(
            RUTA_SHAREPOINT_MTTO,
            f"01. {ano}",
            mes,
            proyecto,
            id_mtto
        )
        
        os.makedirs(
            carpeta_mtto,
            exist_ok=True
        )

        archivos = request.files.getlist(
            "evidencias"
        )

        if not archivos or archivos[0].filename == "":

            return """
            <html>
            <body style="font-family:Arial">

            <h2>Error</h2>

            <p>Debe cargar al menos una evidencia.</p>

            <button onclick="window.history.back()">
            Volver
            </button>

            </body>
            </html>
            """

        contador = 0

        for archivo in archivos:

            if archivo.filename == "":
                continue

            try:

                nombre_seguro = secure_filename(
                    archivo.filename
                )

                if "." not in nombre_seguro:
                    continue

                extension = (
                    nombre_seguro
                    .rsplit(".", 1)[1]
                    .lower()
                )

                if extension not in EXTENSIONES_PERMITIDAS:
                    continue

                contador += 1

                ruta_archivo = os.path.join(
                    carpeta_mtto,
                    nombre_seguro
                )

                if os.path.exists(ruta_archivo):

                    nombre_seguro = (
                        f"{datetime.now().strftime('%H%M%S')}_{nombre_seguro}"
                    )

                    ruta_archivo = os.path.join(
                        carpeta_mtto,
                        nombre_seguro
                    )

                print(
                   "GUARDANDO EN:"
                )

                print(
                    ruta_archivo
                )

                archivo.save(ruta_archivo)

                print(
                    "GUARDADO OK:"
                )

                print(
                    ruta_archivo
                )

            except Exception as e:

                print(
                    "ERROR PROCESANDO ARCHIVO:"
                )

                print(e)

        print(
            f"TOTAL EVIDENCIAS GUARDADAS: {contador}"
        )
        libro = load_workbook(
            ARCHIVO_EXCEL
        )

        hoja = libro["Informe OMCE"]

        print("Ruta a guardar:", carpeta_mtto)

        fila = hoja.max_row + 1

        hoja.cell(fila, 1).value = id_mtto
        hoja.cell(fila, 2).value = fecha
        hoja.cell(fila, 3).value = proyecto
        hoja.cell(fila, 4).value = lider
        hoja.cell(fila, 5).value = tipo
        hoja.cell(fila, 6).value = servicio
        hoja.cell(fila, 7).value = novedades
        hoja.cell(fila, 8).value = pendientes
        hoja.cell(fila, 9).value = fecha_actual
        hoja.cell(fila, 11).value = carpeta_mtto

        fuente = Font(
            name="Calibri",
            size=14
        )

        alineacion = Alignment(
            horizontal="center",
            vertical="center"
        )

        for col in [1, 2, 3, 4, 5, 6, 7, 8, 9, 11]:
            celda = hoja.cell(fila, col)
            celda.font = fuente
            celda.alignment = alineacion

        hoja.row_dimensions[fila].height = 35

        libro.save(
            ARCHIVO_EXCEL
        )

        return f"""
        <!DOCTYPE html>
        <html lang="es">

        <head>

        <meta charset="UTF-8">
        <meta name="viewport" content="width=device-width, initial-scale=1.0">

        <title>Mantenimiento Guardado</title>

        <style>

        body{{
            background:#8BC34A;
            font-family:Arial,sans-serif;
            margin:0;
            padding:20px;
        }}

        .contenedor{{
            max-width:700px;
            margin:30px auto;
            background:white;
            border-radius:15px;
            padding:40px;
            text-align:center;
            box-shadow:0 4px 10px rgba(0,0,0,0.2);
        }}

        .check{{
            font-size:70px;
        }}

        h1{{
            color:#2E7D32;
            margin-top:10px;
        }}

        .id-box{{
            background:#f5f5f5;
            border-left:6px solid #4CAF50;
            padding:20px;
            border-radius:10px;
            margin:30px 0;
        }}

        .id-box p{{
            font-size:32px;
            font-weight:bold;
            margin:10px 0;
            word-break:break-word;
        }}

        .boton{{
            display:block;
            width:100%;
            max-width:450px;
            margin:15px auto;
            padding:18px;
            background:#5D7A1E;
            color:white;
            text-decoration:none;
            border-radius:10px;
            font-size:18px;
            font-weight:bold;
            text-align:center;
        }}

        .boton:hover{{
            background:#4F6819;
        }}

        @media(max-width:768px){{
            .contenedor{{
                padding:25px;
            }}

            .id-box p{{
                font-size:24px;
            }}
        }}

        </style>

        </head>

        <body>

        <div class="contenedor">

        <div class="check">✔️</div>

        <p style="color:#666;font-size:14px;">
        Las evidencias y la información fueron registradas exitosamente.
        </p>

        <h1>Mantenimiento Guardado Correctamente</h1>

        <div class="id-box">
            <h3>ID Generado</h3>
            <p>{id_mtto}</p>
        </div>

        <p>
        Las evidencias fueron almacenadas correctamente.
        </p>

        <a href="/mantenimiento" class="boton">
        📝 Crear Otro Mantenimiento
        </a>

        <a href="/" class="boton">
        🏠 Volver al Menú Principal
        </a>
        </div>

        </body>
        </html>
        """

    return render_template(
        "mantenimiento.html",
        proyectos=proyectos
    )


@app.route("/detalle/<id_mtto>")
def detalle(id_mtto):

    df = pd.read_excel(
        ARCHIVO_EXCEL,
        sheet_name="Informe OMCE",
        engine="openpyxl"
    )

    registro = df[
        df["ID"].astype(str) == str(id_mtto)
    ]

    if registro.empty:
        return "Mantenimiento no encontrado"

    registro = registro.iloc[0]

    archivos = []

    if "RutaLocal" in registro.index:

        ruta = str(registro["RutaLocal"])

        if ruta != "nan" and os.path.exists(ruta):

            archivos = os.listdir(ruta)
    
    print("Ruta encontrada:", ruta)
    print("Existe carpeta:", os.path.exists(ruta))

    archivos = []

    if os.path.exists(ruta):

        print("Archivos encontrados:")

        for archivo in os.listdir(ruta):

            print(archivo)

            archivos.append(archivo)

    return render_template(
        "detalles.html",
        registro=registro,
        archivos=archivos
    )

from flask import send_from_directory

@app.route("/evidencia/<id_mtto>/<archivo>")
def evidencia(id_mtto, archivo):

    df = pd.read_excel(
        ARCHIVO_EXCEL,
        sheet_name="Informe OMCE",
        engine="openpyxl"
    )

    registro = df[
        df["ID"].astype(str) == str(id_mtto)
    ]

    if registro.empty:
        return "No encontrado"

    ruta = str(registro.iloc[0]["RutaLocal"])

    archivo_completo = os.path.join(
        ruta,
        archivo
    )

    print("Archivo solicitado:", archivo_completo)

    if not os.path.exists(archivo_completo):
        return "Archivo no encontrado"

    return send_from_directory(
        os.path.abspath(ruta),
        archivo
    )

@app.route("/comisionamiento", methods=["GET", "POST"])
def comisionamiento():

    df_proyectos = pd.read_excel(
        ARCHIVO_EXCEL,
        sheet_name="PRB",
        engine="openpyxl"
    )

    proyectos = (
        df_proyectos["Proyectos"]
        .dropna()
        .tolist()
    )

    if request.method == "POST":

        proyecto = request.form.get("proyecto")
        lider = request.form.get("lider")
        fecha = request.form.get("fecha")
        tipo = request.form.get("tipocomisionamiento")
        servicio = request.form.get("servicio")
        novedades = request.form.get("novedades")
        pendientes = request.form.get("pendientes")

        if not proyecto or not lider or not fecha or not tipo:

            return """
            <html>
            <body style="font-family:Arial">

            <h2>Error</h2>

            <p>Debe diligenciar todos los campos obligatorios.</p>

            <button onclick="window.history.back()">
            Volver
            </button>

            </body>
            </html>
            """

        fecha_actual = datetime.now()

        id_com = (
            "COM-"
            + fecha_actual.strftime("%Y%m%d")
            + "-"
            + str(random.randint(100000, 999999))
        )

        meses = {
            1: "01. Enero",
            2: "02. Febrero",
            3: "03. Marzo",
            4: "04. Abril",
            5: "05. Mayo",
            6: "06. Junio",
            7: "07. Julio",
            8: "08. Agosto",
            9: "09. Septiembre",
            10: "10. Octubre",
            11: "11. Noviembre",
            12: "12. Diciembre"
        }

        ano = str(fecha_actual.year)
        mes = meses[fecha_actual.month]

        RUTA_SHAREPOINT_COM = r"C:\Users\Usuario\OneDrive - ERCO ENERGIA SAS\AOM - 001_AOM_COMERCIAL&INDUSTRIAL\02_Ejecucion\2.4_AOM_Centro\2.4.4_Mantenimiento\2.4.4.3_Anexos\02. Comisionamientos"
        
        carpeta_com = os.path.join(
            RUTA_SHAREPOINT_COM,
            f"01. {ano}",
            mes,
            proyecto,
            id_com
        )

        os.makedirs(
            carpeta_com,
            exist_ok=True
        )

        archivos = request.files.getlist(
            "evidencias"
        )

        if not archivos or archivos[0].filename == "":

            return """
            <html>
            <body style="font-family:Arial">

            <h2>Error</h2>

            <p>Debe cargar al menos una evidencia.</p>

            <button onclick="window.history.back()">
            Volver
            </button>

            </body>
            </html>
            """

        for archivo in archivos:

            if archivo.filename == "":
                continue

            try:

                nombre_seguro = secure_filename(
                    archivo.filename
                )

                if "." not in nombre_seguro:
                    continue

                extension = (
                    nombre_seguro
                    .rsplit(".", 1)[1]
                    .lower()
                )

                if extension not in EXTENSIONES_PERMITIDAS:
                    continue

                ruta_archivo = os.path.join(
                    carpeta_com,
                    nombre_seguro
                )

                archivo.save(ruta_archivo)

                extension_convertir = {
                    "heic",
                    "heif",
                    "png",
                    "webp",
                    "bmp",
                    "gif",
                    "tif",
                    "tiff"
                }

                if extension in extension_convertir:

                    print("Convirtiendo:", ruta_archivo)

                    ruta_archivo = convertir_a_jpg(
                        ruta_archivo
                    )

                    print(
                        "Convertido a:",
                        ruta_archivo
                    )

            except Exception as e:

                print(
                    "ERROR PROCESANDO ARCHIVO:"
                )

                print(e)

        libro = load_workbook(
            ARCHIVO_EXCEL
        )

        hoja = libro["Informe PRB"]

        fila = hoja.max_row + 1

        hoja.cell(fila, 1).value = id_com
        hoja.cell(fila, 2).value = fecha
        hoja.cell(fila, 3).value = proyecto
        hoja.cell(fila, 4).value = lider
        hoja.cell(fila, 5).value = tipo
        hoja.cell(fila, 6).value = servicio
        hoja.cell(fila, 7).value = novedades
        hoja.cell(fila, 8).value = pendientes
        hoja.cell(fila, 9).value = fecha_actual

        # RutaLocal (columna K)
        hoja.cell(fila, 11).value = carpeta_com

        fuente = Font(
            name="Calibri",
            size=14
        )
        
        alineacion = Alignment(
            horizontal="center",
            vertical="center"
        )
        
        for col in [1, 2, 3, 4, 5, 6, 7, 8, 9, 11]:
            celda = hoja.cell(fila, col)
            celda.font = fuente
            celda.alignment = alineacion
        
            hoja.row_dimensions[fila].height = 35

        libro.save(
            ARCHIVO_EXCEL
        )

        return f"""
        <html>
        <body style="font-family:Arial">

        <h2>Comisionamiento guardado correctamente</h2>

        <h3>ID generado:</h3>

        <h1>{id_com}</h1>

        <p>Las evidencias fueron almacenadas correctamente en SharePoint.:</p>
        
        <b>ID: {id_com}</b>

        <br><br>

        <button onclick="window.location.href='/comisionamiento'">
        Crear otro comisionamiento
        </button>

        <br><br>

        <button onclick="window.location.href='/'">
        Volver al menú
        </button>

        </body>
        </html>
        """

    return render_template(
        "comisionamiento.html",
        proyectos=proyectos
    )

@app.route("/detalle_comisionamiento/<id_com>")
def detalle_comisionamiento(id_com):

    df = pd.read_excel(
        ARCHIVO_EXCEL,
        sheet_name="Informe PRB",
        engine="openpyxl"
    )

    registro = df[
        df["ID"].astype(str) == str(id_com)
    ]

    if registro.empty:
        return "Comisionamiento no encontrado"

    registro = registro.iloc[0]

    ruta = str(registro["RutaLocal"])

    if ruta == "nan":
        return "No existe ruta de evidencias para este comisionamiento"

    print("Ruta encontrada:", ruta)
    print("Existe carpeta:", os.path.exists(ruta))

    archivos = []

    if os.path.exists(ruta):

        print("Archivos encontrados:")

        for archivo in os.listdir(ruta):

            print(archivo)

            archivos.append(archivo)

    return render_template(
        "detalle_comisionamiento.html",
        registro=registro,
        archivos=archivos
    )

from flask import send_from_directory

@app.route("/consultar_comisionamiento")
def consultar_comisionamiento():

    df = pd.read_excel(
        ARCHIVO_EXCEL,
        sheet_name="Informe PRB",
        engine="openpyxl"
    )
    print(df.head())
    print(df.columns.tolist())
    print("Filas:", len(df))

    filtro_id = request.args.get("id", "")
    filtro_proyecto = request.args.get("proyecto", "")
    filtro_lider = request.args.get("lider", "")

    if filtro_id:
        df = df[
            df["ID"].astype(str).str.contains(
                filtro_id,
                case=False,
                na=False
            )
        ]

    if filtro_proyecto:
        df = df[
            df["Proyecto"].astype(str).str.contains(
                filtro_proyecto,
                case=False,
                na=False
            )
        ]

    if filtro_lider:
        df = df[
            df["LiderEncargado"].astype(str).str.contains(
                filtro_lider,
                case=False,
                na=False
            )
        ]

    registros = df.to_dict(orient="records")

    print("Registros encontrados:", len(registros))

    return render_template(
        "consultar_comisionamiento.html",
        registros=registros,
        filtro_id=filtro_id,
        filtro_proyecto=filtro_proyecto,
        filtro_lider=filtro_lider
    )

if __name__ == "__main__":
    app.run(
        host="0.0.0.0",
        port=5000,
        debug=True
    )